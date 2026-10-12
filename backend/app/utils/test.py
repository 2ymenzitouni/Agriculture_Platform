from pathlib import Path
import csv
import tkinter as tk
from tkinter import filedialog

from google import genai
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


# ---------------------------------
# CONFIGURATION
# ---------------------------------

API_KEY = "AQ.Ab8RN6If1PD_9V3sme-T9R9dyTLxvANbjgljojiEXwuOBzVjgw"
MODEL = "gemini-3.1-flash-lite"

client = genai.Client(api_key=API_KEY)

PROMPT = """
You are a document text extraction system.

Extract the original text and data from the provided document.

Rules:
1. Return only the extracted text.
2. Do not explain, summarize, analyze, or interpret anything.
3. Do not convert the content into JSON.
4. Preserve the original wording, numbers, dates, names, units,
   headings, lists, and table values.
5. Preserve the reading order as much as possible.
6. Represent tables as plain text, keeping their headers and rows clear.
7. Do not add information that is not present in the document.
8. Do not correct spelling or grammar.
9. If text is unreadable, write [unreadable].
10. Do not add introductions, conclusions, or comments.
"""


# ---------------------------------
# READ TEXT AND OFFICE DOCUMENTS
# ---------------------------------

def read_document(path: Path) -> str:
    extension = path.suffix.lower()

    if extension in [".txt", ".md"]:
        return path.read_text(
            encoding="utf-8",
            errors="replace"
        )

    elif extension == ".csv":
        with path.open(
            "r",
            encoding="utf-8-sig",
            errors="replace",
            newline=""
        ) as file:
            return "\n".join(
                " | ".join(row)
                for row in csv.reader(file)
            )

    elif extension == ".docx":
        doc = Document(str(path))
        parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        for table in doc.tables:
            parts.append("\n")
            for row in table.rows:
                parts.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        return "\n".join(parts)

    elif extension == ".xlsx":
        workbook = load_workbook(
            filename=str(path),
            read_only=True,
            data_only=True
        )

        parts = []

        for sheet in workbook.worksheets:
            parts.append(f"\n{sheet.title}")

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value) if value is not None else ""
                    for value in row
                ]

                if any(values):
                    parts.append(" | ".join(values))

        workbook.close()
        return "\n".join(parts)

    elif extension == ".pptx":
        presentation = Presentation(str(path))
        parts = []

        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    parts.append(shape.text)

                if shape.has_table:
                    for row in shape.table.rows:
                        parts.append(
                            " | ".join(
                                cell.text for cell in row.cells
                            )
                        )

        return "\n".join(parts)

    else:
        raise ValueError("Unsupported file format.")


# ---------------------------------
# EXTRACT TEXT WITH GEMINI
# ---------------------------------

def extract_file(path: Path) -> str:
    extension = path.suffix.lower()

    image_extensions = {
        ".png", ".jpg", ".jpeg", ".webp",
        ".heic", ".heif"
    }

    # PDF and image files
    if extension == ".pdf" or extension in image_extensions:
        print("Uploading file to Gemini...")

        uploaded_file = client.files.upload(
            file=str(path)
        )

        response = client.models.generate_content(
            model=MODEL,
            contents=[
                PROMPT,
                uploaded_file
            ]
        )

        return response.text or ""

    # Text and Office files
    document_text = read_document(path)

    if not document_text.strip():
        return ""

    return document_text


# ---------------------------------
# SELECT FILE
# ---------------------------------

def select_file():
    root = tk.Tk()
    root.withdraw()

    filename = filedialog.askopenfilename(
        title="Select a document",
        filetypes=[
            (
                "Supported files",
                "*.pdf *.png *.jpg *.jpeg *.webp "
                "*.heic *.heif *.txt *.md *.csv "
                "*.docx *.xlsx *.pptx"
            ),
            ("All files", "*.*")
        ]
    )

    root.destroy()

    return Path(filename) if filename else None


# ---------------------------------
# MAIN PROGRAM
# ---------------------------------

def main():
    path = select_file()

    if path is None:
        print("No file selected.")
        return

    try:
        print(f"\nSelected file: {path.name}")
        print("Extracting text...\n")

        result = extract_file(path)

        if not result.strip():
            print("No text was extracted.")
            return

        print("=" * 60)
        print("EXTRACTED TEXT")
        print("=" * 60)
        print(result)

        output_path = path.with_name(
            path.stem + "_extracted.txt"
        )

        output_path.write_text(
            result,
            encoding="utf-8"
        )

        print(f"\nText saved to: {output_path}")

    except Exception as error:
        print(f"\nError: {error}")


if __name__ == "__main__":
    main()
